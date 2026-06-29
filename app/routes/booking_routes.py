from fastapi import APIRouter
from app.schemas.booking_schema import BookingCreate
from app.models.booking import Booking
from app.database.db import SessionLocal
from app.models.room import Room
from fastapi import APIRouter, HTTPException
from sqlalchemy import and_, or_
from app.models.customers import Customer
from datetime import date  
from fastapi import Query
from fastapi.responses import FileResponse
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import inch
import os
from fastapi import Depends
from app.auth.oauth2 import get_current_admin, get_current_customer

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from tempfile import NamedTemporaryFile

router = APIRouter()


@router.get("/bookings")
def get_bookings(
    current_user=Depends(get_current_admin)
):

    db = SessionLocal()

    bookings = db.query(Booking).order_by(
        Booking.booking_id
    ).all()

    result = []

    for booking in bookings:

        result.append(
            {
                "booking_id": booking.booking_id,
                "customer_id": booking.customer_id,
                "room_id": booking.room_id,
                "check_in": booking.check_in,
                "check_out": booking.check_out,
                "booking_status": booking.booking_status
            }
        )

    db.close()

    return result


@router.post("/bookings")
def create_booking(
    booking: BookingCreate,
    current_user=Depends(get_current_customer)
):
    db = SessionLocal()

    # Check-Out must be after Check-In
    if booking.check_out <= booking.check_in:

        db.close()

        raise HTTPException(
            status_code=400,
            detail="Check-out date must be after Check-in date"
        )
    if booking.check_in < date.today():

            db.close()

            raise HTTPException(
                status_code=400,
                detail="Check-in date cannot be in the past"
        )

    # Check if room already booked for selected dates
    existing_booking = db.query(Booking).filter(
        Booking.room_id == booking.room_id,
        Booking.check_in < booking.check_out,
        Booking.check_out > booking.check_in
    ).first()

    if existing_booking:

        db.close()

        raise HTTPException(
            status_code=400,
            detail="Room already booked for selected dates"
        )

    new_booking = Booking(
        customer_id=booking.customer_id,
        room_id=booking.room_id,
        check_in=booking.check_in,
        check_out=booking.check_out,
        booking_status=booking.booking_status
    )

    db.add(new_booking)
    db.commit()
    db.refresh(new_booking)

    room = db.query(Room).filter(
        Room.room_id == booking.room_id
    ).first()

    if room:
        room.status = "occupied"
        db.commit()

    booking_id = new_booking.booking_id

    db.close()

    return {
        "message": "Booking Created Successfully",
        "booking_id": booking_id
    }
@router.put("/bookings/{booking_id}")
def update_booking(
    booking_id:int,
    booking:BookingCreate,
    current_user=Depends(get_current_admin)
):
    db = SessionLocal()

    existing_booking = db.query(Booking).filter(
        Booking.booking_id == booking_id
    ).first()

    if not existing_booking:

        db.close()

        return {
            "message": "Booking Not Found"
        }

    existing_booking.customer_id = booking.customer_id
    existing_booking.room_id = booking.room_id
    existing_booking.check_in = booking.check_in
    existing_booking.check_out = booking.check_out
    existing_booking.booking_status = booking.booking_status

    db.commit()

    db.close()

    return {
        "message": f"Booking {booking_id} Updated Successfully"
    }
@router.delete("/bookings/{booking_id}")
def delete_booking(
    booking_id: int,
    current_user=Depends(get_current_admin)
):

    db = SessionLocal()

    booking = db.query(Booking).filter(
        Booking.booking_id == booking_id
    ).first()

    if not booking:
        db.close()
        return {
            "message": "Booking Not Found"
        }

    room = db.query(Room).filter(
        Room.room_id == booking.room_id
    ).first()

    if room:
        room.status = "available"

    db.delete(booking)

    db.commit()

    db.close()

    return {
        "message": "Booking Deleted Successfully"
    }


@router.get("/my-bookings/{customer_id}")
def get_my_bookings(
    customer_id:int,
    current_user=Depends(get_current_customer)
):

    db = SessionLocal()

    bookings = db.query(
        Booking,
        Room
    ).join(
        Room,
        Booking.room_id == Room.room_id
    ).filter(
        Booking.customer_id == customer_id
    ).all()

    result = []

    for booking, room in bookings:

        status = booking.booking_status

        if (
            status != "cancelled"
            and booking.check_out < date.today()
        ):
            status = "completed"

        result.append(
            {
                "booking_id": booking.booking_id,
                "room_number": room.room_number,
                "room_type": room.room_type,
                "price": float(room.price),
                "check_in": booking.check_in,
                "check_out": booking.check_out,
                "booking_status": status
            }
        )

    db.close()

    return result
@router.get("/dashboard-stats")
def dashboard_stats(
    current_user=Depends(get_current_admin)
):

    db = SessionLocal()

    total_rooms = db.query(Room).count()

    total_bookings = db.query(Booking).count()

    total_customers = db.query(Customer).count()

    revenue = 0

    bookings = db.query(
        Booking,
        Room
    ).join(
        Room,
        Booking.room_id == Room.room_id
    ).all()

    for booking, room in bookings:

        if booking.booking_status != "cancelled":

            revenue += float(room.price)

    available_rooms = db.query(Room).filter(
        Room.status == "available"
    ).count()

    occupied_rooms = db.query(Room).filter(
        Room.status == "occupied"
    ).count()

    cancelled_bookings = db.query(Booking).filter(
        Booking.booking_status == "cancelled"
    ).count()

    pending_bookings = db.query(Booking).filter(
        Booking.booking_status == "booked"
    ).count()

    print("Available Rooms =", available_rooms)
    print("Occupied Rooms =", occupied_rooms)
    print("Cancelled Bookings =", cancelled_bookings)
    print("Pending Bookings =", pending_bookings)

    db.close()

    return {
        "total_rooms": total_rooms,
        "total_bookings": total_bookings,
        "total_customers": total_customers,
        "total_revenue": revenue,

        "available_rooms": available_rooms,
        "occupied_rooms": occupied_rooms,
        "cancelled_bookings": cancelled_bookings,
        "pending_bookings": pending_bookings
    }
@router.get("/recent-bookings")
def recent_bookings(
    current_user=Depends(get_current_admin)
):

    db = SessionLocal()

    bookings = db.query(
        Booking,
        Customer,
        Room
    ).join(
        Customer,
        Booking.customer_id == Customer.customer_id
    ).join(
        Room,
        Booking.room_id == Room.room_id
    ).order_by(
        Booking.booking_id.desc()
    ).limit(5).all()

    result = []

    for booking, customer, room in bookings:

        result.append(
            {
                "booking_id": booking.booking_id,
                "customer_name": customer.name,
                "room_number": room.room_number,
                "check_in": str(booking.check_in),
                "check_out": str(booking.check_out),
                "status": booking.booking_status,
                "amount": float(room.price)
            }
        )

    db.close()

    return result
@router.get("/admin-bookings-data")
def admin_bookings_data(
    selected_date: date | None = Query(None),
    history: bool = Query(False),
    current_user=Depends(get_current_admin)
):

    db = SessionLocal()

    query = db.query(
        Booking,
        Customer,
        Room
    ).join(
        Customer,
        Booking.customer_id == Customer.customer_id
    ).join(
        Room,
        Booking.room_id == Room.room_id
    )

    if not history and selected_date:

        query = query.filter(
    Booking.check_in == selected_date
)

    bookings = query.all()

    print("Selected Date :", selected_date)
    print("History :", history)
    print("Bookings Found :", len(bookings))

    result = []

    for booking, customer, room in bookings:

        result.append(
            {
                "booking_id": booking.booking_id,
                "customer_id": customer.customer_id,
                "customer_name": customer.name,
                "room_id": room.room_id,
                "room_number": room.room_number,
                "check_in": booking.check_in,
                "check_out": booking.check_out,
                "booking_status": booking.booking_status
            }
        )

    db.close()

    return result
@router.put("/checkout/{booking_id}")
def checkout_booking(
    booking_id:int,
    current_user=Depends(get_current_admin)
):

    db = SessionLocal()

    booking = db.query(Booking).filter(
        Booking.booking_id == booking_id
    ).first()

    if not booking:
        db.close()
        return {"message": "Booking Not Found"}

    booking.booking_status = "completed"

    room = db.query(Room).filter(
        Room.room_id == booking.room_id
    ).first()

    if room:
        room.status = "available"

    db.commit()
    db.close()

    return {
        "message": "Checkout Successful"
    }
@router.get("/export-report")
def export_report(
    start_date: date = Query(...),
    end_date: date = Query(...),
    current_user=Depends(get_current_admin)
):

    db = SessionLocal()

    bookings = get_report_bookings(
    db,
    start_date,
    end_date
)

    total_revenue = 0

    table_data = [
        [
            "Booking ID",
            "Customer",
            "Room",
            "Check In",
            "Check Out",
            "Status"
        ]
    ]

    for booking, customer, room in bookings:

        if booking.booking_status != "cancelled":
            total_revenue += float(room.price)

        table_data.append([
            str(booking.booking_id),
            customer.name,
            str(room.room_number),
            str(booking.check_in),
            str(booking.check_out),
            booking.booking_status
        ])

    pdf_name = f"Hotel_Report_{start_date}_to_{end_date}.pdf"

    doc = SimpleDocTemplate(pdf_name)

    styles = getSampleStyleSheet()

    elements = []

    elements.append(
        Paragraph(
            "<b><font size=18>HOTEL MANAGEMENT REPORT</font></b>",
            styles["Title"]
        )
    )

    elements.append(
        Paragraph(
            f"<b>Period :</b> {start_date} to {end_date}",
            styles["Normal"]
        )
    )

    elements.append(
        Paragraph(
            f"<b>Total Bookings :</b> {len(bookings)}",
            styles["Normal"]
        )
    )

    elements.append(
        Paragraph(
            f"<b>Total Revenue :</b> ₹{total_revenue}",
            styles["Normal"]
        )
    )

    elements.append(
        Paragraph("<br/><br/>", styles["Normal"])
    )

    table = Table(table_data)

    table.setStyle(TableStyle([

        ("BACKGROUND",(0,0),(-1,0),colors.darkblue),

        ("TEXTCOLOR",(0,0),(-1,0),colors.white),

        ("GRID",(0,0),(-1,-1),1,colors.black),

        ("BACKGROUND",(0,1),(-1,-1),colors.beige),

        ("ALIGN",(0,0),(-1,-1),"CENTER"),

        ("BOTTOMPADDING",(0,0),(-1,0),10),

    ]))

    elements.append(table)

    doc.build(elements)

    db.close()

    return FileResponse(
        pdf_name,
        media_type="application/pdf",
        filename=pdf_name
    )
@router.get("/export-excel")
def export_excel(
    start_date: date = Query(...),
    end_date: date = Query(...),
    current_user=Depends(get_current_admin)
):
    db = SessionLocal()

    bookings = get_report_bookings(
    db,
    start_date,
    end_date
)
    wb = Workbook()

    ws = wb.active

    ws.title = "Hotel Report"

    # -------------------------
    # Title
    # -------------------------

    ws["A1"] = "HOTEL MANAGEMENT REPORT"

    ws["A1"].font = Font(
        bold=True,
        size=16
    )

    ws["A3"] = f"Period : {start_date} to {end_date}"

    # -------------------------
    # Header
    # -------------------------

    headers = [

        "Booking ID",

        "Customer",

        "Room",

        "Check In",

        "Check Out",

        "Status",

        "Amount"

    ]

    header_fill = PatternFill(

        start_color="1F4E78",

        end_color="1F4E78",

        fill_type="solid"

    )

    row = 5

    for col, header in enumerate(headers, start=1):

        cell = ws.cell(row=row, column=col)

        cell.value = header

        cell.font = Font(
            bold=True,
            color="FFFFFF"
        )

        cell.fill = header_fill

    # -------------------------
    # Data
    # -------------------------

    total_revenue = 0

    row = 6

    for booking, customer, room in bookings:

        if booking.booking_status != "cancelled":

            total_revenue += float(room.price)

        ws.cell(row=row, column=1).value = booking.booking_id

        ws.cell(row=row, column=2).value = customer.name

        ws.cell(row=row, column=3).value = room.room_number

        ws.cell(row=row, column=4).value = str(booking.check_in)

        ws.cell(row=row, column=5).value = str(booking.check_out)

        ws.cell(row=row, column=6).value = booking.booking_status

        ws.cell(row=row, column=7).value = float(room.price)

        row += 1

    # -------------------------
    # Summary
    # -------------------------

    ws.cell(row=row + 2, column=1).value = "Total Bookings"

    ws.cell(row=row + 2, column=2).value = len(bookings)

    ws.cell(row=row + 3, column=1).value = "Total Revenue"

    ws.cell(row=row + 3, column=2).value = total_revenue

    temp_file = NamedTemporaryFile(
    delete=False,
    suffix=".xlsx"
)

    temp_file.close()

    wb.save(temp_file.name)

    db.close()

    return FileResponse(

        temp_file.name,

        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",

        filename=f"Hotel_Report_{start_date}_to_{end_date}.xlsx"

    )

def get_report_bookings(db, start_date, end_date):

    return db.query(
        Booking,
        Customer,
        Room
    ).join(
        Customer,
        Booking.customer_id == Customer.customer_id
    ).join(
        Room,
        Booking.room_id == Room.room_id
    ).filter(

        # Overlap Logic
        Booking.check_in <= end_date,
        Booking.check_out >= start_date

    ).all()
@router.get("/reports-data")
def reports_data(
    selected_date: date | None = Query(None),
    current_user=Depends(get_current_admin)
):

    db = SessionLocal()

    # ---------------- Rooms ----------------

    total_rooms = db.query(Room).count()

    total_customers = db.query(Customer).count()

    room_query = db.query(Room)

    if selected_date:

        occupied_room_ids = db.query(
            Booking.room_id
        ).filter(
            Booking.check_in <= selected_date,
            Booking.check_out > selected_date
        ).all()

        occupied_room_ids = [r[0] for r in occupied_room_ids]

        occupied_rooms = len(occupied_room_ids)
        available_rooms = total_rooms - occupied_rooms

    else:

        available_rooms = db.query(Room).filter(
            Room.status == "available"
        ).count()

        occupied_rooms = db.query(Room).filter(
            Room.status == "occupied"
        ).count()

    # ---------------- Bookings ----------------

    booking_query = db.query(
        Booking,
        Customer,
        Room
    ).join(
        Customer,
        Booking.customer_id == Customer.customer_id
    ).join(
        Room,
        Booking.room_id == Room.room_id
    )

    if selected_date:

        booking_query = booking_query.filter(
        Booking.check_in <= selected_date,
        Booking.check_out >= selected_date
    )
    bookings = booking_query.all()

    total_bookings = len(bookings)

    booked = 0
    completed = 0
    cancelled = 0
    revenue = 0

    recent = []

    for booking, customer, room in bookings:

        if booking.booking_status == "booked":
            booked += 1

        elif booking.booking_status == "completed":
            completed += 1

        elif booking.booking_status == "cancelled":
            cancelled += 1

        if booking.booking_status != "cancelled":
            revenue += float(room.price)

        recent.append({
            "booking_id": booking.booking_id,
            "customer_name": customer.name,
            "room_number": room.room_number,
            "check_in": str(booking.check_in),
            "check_out": str(booking.check_out),
            "status": booking.booking_status
        })

    recent = recent[::-1][:5]

    db.close()

    return {
        "total_rooms": total_rooms,
        "total_customers": total_customers,
        "total_bookings": total_bookings,
        "total_revenue": revenue,

        "available_rooms": available_rooms,
        "occupied_rooms": occupied_rooms,

        "booked": booked,
        "completed": completed,
        "cancelled": cancelled,

        "recent_bookings": recent
    }